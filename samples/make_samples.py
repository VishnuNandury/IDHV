"""Generate synthetic sample documents for the IDH Vault demo.

Patient: Ananya Sharma (DHID DH-8923-4102-9981, F, DOB 1968-04-12).
All data is synthetic. Run:  python samples/make_samples.py
Needs: pillow, reportlab, python-docx, openpyxl (and ffmpeg on PATH for the mp4).
"""
import os
import random
import shutil
import subprocess
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont

OUT = Path(__file__).resolve().parent
FONTS = Path("C:/Windows/Fonts")
random.seed(7)

PATIENT = "Ananya Sharma"
DHID = "DH-8923-4102-9981"
RS = "\u20b9"


def font(names, size):
    for n in names:
        for base in (FONTS, Path("/usr/share/fonts/truetype/dejavu")):
            p = base / n
            if p.exists():
                return ImageFont.truetype(str(p), size)
    return ImageFont.load_default()


def hand(size):
    f = font(["Inkfree.ttf", "segoesc.ttf", "LHANDW.TTF"], size)
    if not isinstance(f, ImageFont.FreeTypeFont):
        print("WARNING: no handwriting font found; falling back to default bitmap font")
    return f


def sans(size, bold=False):
    return font(["arialbd.ttf" if bold else "arial.ttf", "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"], size)


def mono(size):
    return font(["cour.ttf", "consola.ttf", "DejaVuSansMono.ttf"], size)


def add_noise(img, amount=12):
    px = img.load()
    w, h = img.size
    for _ in range(w * h // 18):
        x, y = random.randrange(w), random.randrange(h)
        r, g, b = px[x, y]
        d = random.randint(-amount, amount)
        px[x, y] = (max(0, min(255, r + d)), max(0, min(255, g + d)), max(0, min(255, b + d)))
    return img


# ---------------------------------------------------------------- prescription
def make_prescription():
    W, H = 1100, 1450
    img = Image.new("RGB", (W, H), (250, 247, 238))
    d = ImageDraw.Draw(img)
    ink = (28, 42, 150)
    printed = (60, 60, 60)
    d.text((60, 50), "Dr. R. MENON, MBBS MD (Gen. Med.)", font=sans(34, True), fill=printed)
    d.text((60, 95), "Menon Clinic, Kothrud, Pune - 411038   Reg. No. MMC-48213", font=sans(22), fill=printed)
    d.text((60, 125), "Ph: 020-2543 8800", font=sans(22), fill=printed)
    d.line((50, 165, W - 50, 165), fill=printed, width=3)

    h = hand(40)
    hs = hand(34)
    d.text((70, 200), "Date: 15/09/2026", font=h, fill=ink)
    d.text((600, 200), "Pt: Ananya Sharma  58/F", font=h, fill=ink)
    d.text((70, 262), "C/o  T2DM (2019), HTN (2021)", font=h, fill=ink)
    d.text((70, 322), "BP 138/86   Wt 68 kg   RBS 176", font=h, fill=ink)
    d.text((70, 420), "Rx", font=hand(80), fill=ink)
    lines = [
        "1) Tab Glycomet 500    1 - 0 - 1  (BD)",
        "2) Tab Januvia 100     1 - 0 - 0  (OD)",
        "3) Tab Telma 40        1 - 0 - 0  (OD)",
        "4) Tab Atorva 20       0 - 0 - 1  (HS)",
        "5) Tab Ecosprin 75     1 - 0 - 0  (OD)",
        "    x 30 days",
    ]
    y = 500
    for ln in lines:
        d.text((110, y), ln, font=h, fill=ink)
        y += 70
    d.text((70, y + 20), "Adv: HbA1c, lipid profile, creatinine", font=h, fill=ink)
    d.text((70, y + 90), "Low salt diet, walk 30 min daily", font=h, fill=ink)
    d.text((70, y + 170), "Review after 3 months  ~ 15/12/26", font=hs, fill=ink)
    d.text((70, y + 260), f"Fee: {RS}700 /-  (paid)", font=h, fill=ink)
    d.text((740, H - 200), "R. Menon", font=hand(64), fill=ink)
    d.line((700, H - 130, W - 80, H - 130), fill=ink, width=2)
    img = add_noise(img)
    img = img.rotate(-1.6, resample=Image.BICUBIC, fillcolor=(250, 247, 238), expand=False)
    img = img.filter(ImageFilter.GaussianBlur(0.6))
    return img


# ------------------------------------------------------------------ lab report
LAB_P1 = [
    ("HbA1c (Glycated Hb)", "7.0", "%", "4.0 - 5.6", "H"),
    ("Fasting Blood Sugar (FBS)", "128", "mg/dL", "70 - 99", "H"),
    ("Post-prandial Blood Sugar (PPBS)", "182", "mg/dL", "70 - 139", "H"),
]
LAB_P2 = [
    ("Total Cholesterol", "196", "mg/dL", "< 200", ""),
    ("Triglycerides", "168", "mg/dL", "< 150", "H"),
    ("HDL Cholesterol", "44", "mg/dL", "> 50", "L"),
    ("LDL Cholesterol", "118", "mg/dL", "< 100", "H"),
    ("Serum Creatinine", "0.9", "mg/dL", "0.5 - 1.1", ""),
    ("TSH", "3.2", "uIU/mL", "0.4 - 4.2", ""),
]


def make_lab_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ss = getSampleStyleSheet()

    def header():
        return [
            Paragraph("<b>Sunrise Diagnostics, Pune</b>", ss["Title"]),
            Paragraph("NABL-style demo lab | 14 FC Road, Pune 411004 | Ph 020-2567 1100", ss["Normal"]),
            Spacer(1, 6 * mm),
            Table(
                [
                    ["Patient", f"{PATIENT} (F, 58 y)", "Sample ID", "SD-260912-0417"],
                    ["DHID", DHID, "Collected", "12/09/2026 08:10"],
                    ["Referred by", "Dr. R. Menon", "Reported", "12/09/2026 16:45"],
                ],
                colWidths=[28 * mm, 55 * mm, 28 * mm, 55 * mm],
                style=TableStyle([("FONTSIZE", (0, 0), (-1, -1), 9), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey)]),
            ),
            Spacer(1, 8 * mm),
        ]

    def result_table(title, rows):
        data = [["Test", "Result", "Unit", "Reference range", "Flag"]] + [list(r) for r in rows]
        t = Table(data, colWidths=[62 * mm, 22 * mm, 22 * mm, 38 * mm, 15 * mm])
        st = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d9e8f5")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
        ]
        for i, r in enumerate(rows, start=1):
            if r[4]:
                st.append(("TEXTCOLOR", (1, i), (1, i), colors.red))
                st.append(("TEXTCOLOR", (4, i), (4, i), colors.red))
        t.setStyle(TableStyle(st))
        return [Paragraph(f"<b>{title}</b>", ss["Heading2"]), t]

    story = header()
    story += result_table("Diabetes Panel", LAB_P1)
    story += [Spacer(1, 6 * mm), Paragraph("Comment: HbA1c 7.0 % indicates fair glycaemic control on current therapy. H = high, L = low.", ss["Normal"])]
    story += [PageBreak()] + header()
    story += result_table("Lipid Profile, Renal and Thyroid", LAB_P2)
    story += [Spacer(1, 6 * mm), Paragraph("Method: enzymatic / CLIA. Fasting 10-12 h. Electronically signed: Dr. S. Kulkarni, MD (Pathology).", ss["Normal"])]
    SimpleDocTemplate(str(OUT / "lab_report.pdf"), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, title="Lab Report", invariant=1).build(story)


# ---------------------------------------------------------------- discharge docx
def make_discharge_docx():
    from docx import Document

    doc = Document()
    doc.add_heading("DISCHARGE SUMMARY", 0)
    doc.add_paragraph("Ruby Hall Clinic (demo), Pune")
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for a, b in [
        ("Patient", f"{PATIENT}, 57 y / F"),
        ("DHID", DHID),
        ("Date of admission", "2025-11-08"),
        ("Date of discharge", "2025-11-10"),
        ("Treating doctor", "Dr. R. Menon, MD (Gen. Med.)"),
        ("Diagnosis", "Hypoglycaemia episode on sulfonylurea therapy; Type 2 diabetes mellitus; Hypertension"),
    ]:
        r = t.add_row().cells
        r[0].text, r[1].text = a, b
    doc.add_heading("Presenting complaints", 2)
    doc.add_paragraph(
        "Sudden sweating, giddiness and confusion at 06:30 after a missed breakfast. "
        "Random blood sugar at arrival 48 mg/dL. BP 128/78 mmHg, pulse 88/min."
    )
    doc.add_heading("Course in hospital", 2)
    doc.add_paragraph(
        "Given IV 25% dextrose followed by 5% dextrose infusion; sugars stabilised (FBS 112 mg/dL on day 2). "
        "ECG normal. Glimepiride STOPPED because of sulfonylurea hypoglycaemia; Sitagliptin (Januvia) 100 mg OD added to Metformin. "
        "Diet counselling done."
    )
    doc.add_heading("Investigations", 2)
    t2 = doc.add_table(rows=1, cols=3)
    t2.style = "Table Grid"
    for c, h in zip(t2.rows[0].cells, ["Test", "Result", "Reference"]):
        c.text = h
    for row in [("HbA1c", "7.4 %", "4.0 - 5.6"), ("Random glucose (admission)", "48 mg/dL", "70 - 140"),
                ("FBS (day 2)", "112 mg/dL", "70 - 99"), ("Serum creatinine", "0.9 mg/dL", "0.5 - 1.1")]:
        cells = t2.add_row().cells
        for c, v in zip(cells, row):
            c.text = v
    doc.add_heading("Medications on discharge", 2)
    t3 = doc.add_table(rows=1, cols=4)
    t3.style = "Table Grid"
    for c, h in zip(t3.rows[0].cells, ["Drug", "Dose", "Frequency", "Duration"]):
        c.text = h
    for row in [("Tab Glycomet (Metformin) 500 mg", "1 tab", "BD (with meals)", "ongoing"),
                ("Tab Januvia (Sitagliptin) 100 mg", "1 tab", "OD", "ongoing"),
                ("Tab Telma (Telmisartan) 40 mg", "1 tab", "OD (morning)", "ongoing"),
                ("Tab Atorva (Atorvastatin) 20 mg", "1 tab", "HS (night)", "ongoing"),
                ("Tab Ecosprin (Aspirin) 75 mg", "1 tab", "OD", "ongoing")]:
        cells = t3.add_row().cells
        for c, v in zip(cells, row):
            c.text = v
    doc.add_heading("Follow-up", 2)
    doc.add_paragraph(f"Review in OPD after 2 weeks with fasting and post-meal sugars. Total hospital bill: {RS}18,450.")
    doc.add_paragraph("Dr. R. Menon")
    import datetime as _dt
    cp = doc.core_properties
    cp.created = cp.modified = _dt.datetime(2025, 11, 10, 12, 0, 0)
    cp.author = cp.last_modified_by = "IDH Vault samples"
    doc.save(OUT / "discharge_summary.docx")


# ------------------------------------------------------------------ xlsx
def make_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "HbA1c"
    ws.append(["Date", "Test", "Value", "Unit", "Reference"])
    for r in [("2024-01-18", 7.9), ("2024-04-20", 7.7), ("2024-07-22", 7.5), ("2024-10-19", 7.4),
              ("2025-01-25", 7.3), ("2025-04-26", 7.2), ("2025-07-19", 7.1), ("2025-11-09", 7.4),
              ("2026-03-14", 7.2), ("2026-06-13", 7.1), ("2026-09-12", 7.0)]:
        ws.append([r[0], "HbA1c", r[1], "%", "4.0 - 5.6"])
    ws2 = wb.create_sheet("Lipids")
    ws2.append(["Date", "Total Cholesterol", "LDL", "HDL", "Triglycerides", "Unit"])
    for r in [("2024-02-10", 218, 142, 42, 190), ("2024-08-17", 205, 130, 43, 182), ("2025-02-15", 201, 124, 44, 176),
              ("2025-08-16", 198, 121, 44, 171), ("2026-03-14", 197, 119, 44, 170), ("2026-09-12", 196, 118, 44, 168)]:
        ws2.append(list(r) + ["mg/dL"])
    for s in (ws, ws2):
        for c in s[1]:
            c.font = Font(bold=True)
        for col in "ABCDEF":
            s.column_dimensions[col].width = 18
    import datetime as _dt
    wb.properties.created = wb.properties.modified = _dt.datetime(2026, 9, 12, 12, 0, 0)
    wb.properties.creator = wb.properties.lastModifiedBy = "IDH Vault samples"
    wb.save(OUT / "lab_history.xlsx")


# ------------------------------------------------------------------ pharmacy bill
def make_pharmacy_bill():
    W, H = 700, 1000
    img = Image.new("RGB", (W, H), (255, 255, 252))
    d = ImageDraw.Draw(img)
    f, fb = mono(24), font(["courbd.ttf", "cour.ttf"], 26)
    y = 30
    for t, bold in [("   WELLCARE MEDICALS", True), ("   Shop 4, Karve Rd, Pune 411004", False),
                    ("   GSTIN 27ABCDE1234F1Z5  DL 20B-4471", False), ("-" * 42, False),
                    ("Bill No: WM/26/09871  Date: 16/09/2026", False), (f"Patient: {PATIENT}", False),
                    ("Dr: R. Menon", False), ("-" * 42, False),
                    ("Item              Qty   MRP    Total", True)]:
        d.text((25, y), t, font=fb if bold else f, fill=(20, 20, 20))
        y += 36
    items = [("Glycomet 500 (20s)", 3, 24.50), ("Januvia 100 (7s)", 4, 215.00),
             ("Telma 40 (15s)", 2, 205.00), ("Atorva 20 (15s)", 2, 178.25),
             ("Ecosprin 75 (14s)", 2, 5.50)]
    total = 0.0
    for name, q, mrp in items:
        t = q * mrp
        total += t
        d.text((25, y), f"{name[:18]:<18} {q:>3} {mrp:>7.2f} {t:>8.2f}", font=f, fill=(20, 20, 20))
        y += 36
    disc = round(total * 0.1, 2)
    d.text((25, y), "-" * 42, font=f, fill=(20, 20, 20)); y += 36
    for lbl, v in [("Sub total", total), ("Discount 10%", -disc), ("GST incl.", 0.0), (f"NET PAYABLE {RS}", total - disc)]:
        d.text((25, y), f"{lbl:<24}{v:>14.2f}", font=fb if "NET" in lbl else f, fill=(20, 20, 20))
        y += 36
    d.text((25, y + 20), "   Paid by UPI. Thank you. Get well soon!", font=f, fill=(20, 20, 20))
    img = add_noise(img, 8).rotate(0.7, resample=Image.BICUBIC, fillcolor=(255, 255, 252))
    img.save(OUT / "pharmacy_bill.jpg", quality=88)
    return total - disc


# ------------------------------------------------------------------ report image + video
def make_report_image():
    W, H = 1280, 720
    img = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, W, 90), fill=(30, 80, 140))
    d.text((40, 24), "Sunrise Diagnostics, Pune - Lab Report", font=sans(40, True), fill="white")
    d.text((40, 115), f"Patient: {PATIENT}   DHID: {DHID}   Date: 12/09/2026", font=sans(26), fill=(30, 30, 30))
    y = 190
    d.text((40, y), "Test                          Result     Unit      Reference    Flag", font=mono(28), fill=(0, 0, 0))
    for n, v, u, r, fl in LAB_P1 + LAB_P2[:2]:
        y += 62
        d.text((40, y), f"{n[:28]:<28} {v:>7}   {u:<8} {r:<12} {fl}", font=mono(28), fill=(180, 0, 0) if fl else (0, 0, 0))
    return img


def make_video(rx_img):
    if not shutil.which("ffmpeg"):
        print("ffmpeg not found; skipping followup_video.mp4")
        return
    with tempfile.TemporaryDirectory() as td:
        a = Path(td) / "a.png"
        b = Path(td) / "b.png"
        make_report_image().save(a)
        rx = rx_img.copy().resize((int(rx_img.width * 720 / rx_img.height), 720))
        canvas = Image.new("RGB", (1280, 720), (240, 240, 240))
        canvas.paste(rx, ((1280 - rx.width) // 2, 0))
        canvas.save(b)
        out = OUT / "followup_video.mp4"
        cmd = ["ffmpeg", "-y", "-loglevel", "error", "-loop", "1", "-t", "4", "-i", str(a), "-loop", "1", "-t", "4", "-i", str(b),
               "-filter_complex", "[0:v][1:v]concat=n=2:v=1:a=0,format=yuv420p[v]", "-map", "[v]", "-r", "10",
               "-c:v", "libx264", "-movflags", "+faststart", str(out)]
        subprocess.run(cmd, check=True)


def _fix_zip(path):
    """Rewrite a zip (docx/xlsx) with fixed entry timestamps for byte-stable output."""
    import zipfile
    tmp = str(path) + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in sorted(zin.infolist(), key=lambda i: i.filename):
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 9, 12, 12, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            data = zin.read(info.filename)
            if info.filename == "docProps/core.xml":
                import re
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>2026-09-12T12:00:00Z", data)
            zout.writestr(zi, data)
    os.replace(tmp, path)


def main():
    rx = make_prescription()
    rx.save(OUT / "prescription_handwritten.jpg", quality=90)
    # image-only PDF (no text layer) to exercise OCR
    rx.save(OUT / "scanned_prescription.pdf", "PDF", resolution=130.0)
    make_lab_pdf()
    make_discharge_docx()
    make_xlsx()
    _fix_zip(OUT / "discharge_summary.docx")
    _fix_zip(OUT / "lab_history.xlsx")
    net = make_pharmacy_bill()
    make_video(rx)
    print("done; pharmacy net", net)
    for p in sorted(OUT.iterdir()):
        print(p.name, p.stat().st_size)


if __name__ == "__main__":
    main()
