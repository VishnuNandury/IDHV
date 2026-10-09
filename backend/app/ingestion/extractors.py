"""Multi-format content extraction -> ``ExtractedContent`` (spec section 5, step 2).

Every result has ``meta["kind"]``. PDFs also carry ``meta["pdf_bytes"]`` (so the Claude provider can send the
original document) and ``meta["page_count"]``. Images are PNG bytes, long edge <= ``MAX_IMAGE_EDGE``.

Hostile-input guards: Pillow pixel limit (decompression bombs, warnings promoted to errors), OOXML zip
pre-check (total size + compression ratio), text output cap for docx/xlsx/xls, and ffmpeg/ffprobe restricted
to the ``file`` protocol.
"""
from __future__ import annotations

import io
import json
import logging
import re
import shutil
import subprocess
import tempfile
import threading
import warnings
import zipfile
from pathlib import Path

from PIL import Image

from app.schemas import ExtractedContent

_PDFIUM_LOCK = threading.Lock()  # PDFium is not thread-safe: serialise every pypdfium2 use

log = logging.getLogger(__name__)

MAX_IMAGE_EDGE = 1568
MIN_PAGE_TEXT = 40  # PDF pages with less text than this are rendered to images
MAX_RENDERED_PAGES = 20
MAX_VIDEO_FRAMES = 6
FFMPEG_TIMEOUT = 120
MAX_TEXT_CHARS = 500_000  # cap on text produced from docx/xlsx/xls
MAX_ZIP_UNCOMPRESSED = 200 * 1024 * 1024  # docx/xlsx: total uncompressed size
MAX_ZIP_RATIO = 100  # docx/xlsx: per-entry compression ratio ...
_ZIP_RATIO_MIN_SIZE = 1024 * 1024  # ... checked for entries larger than 1 MB uncompressed (tiny XML can't bomb)
IMAGE_FORMATS = ["JPEG", "PNG", "WEBP"]
_FFMPEG_SAFE = ["-protocol_whitelist", "file"]

MAX_SHEET_ROWS = 50_000  # rows read per spreadsheet sheet, empty rows included

# Decompression-bomb guard for every Pillow decode in the process (images, video frames). 100 MP lets 48-50 MP
# phone photos through; Pillow raises above 2x this and we promote its warning (above 1x) to an error.
Image.MAX_IMAGE_PIXELS = 100_000_000


class ExtractionError(Exception):
    """The file could not be read (corrupt, unsupported, hostile, or a required tool such as ffmpeg is missing)."""


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _cell(v) -> str:
    if v is None:
        return ""
    if hasattr(v, "isoformat"):
        s = v.isoformat()
        return s[:10] if s.endswith("T00:00:00") else s
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).replace("\t", " ").replace("\n", " ").strip()


class _TextBudget:
    """Collects output lines until ``MAX_TEXT_CHARS`` is reached, then appends one truncation note."""

    def __init__(self) -> None:
        self.lines: list[str] = []
        self.used = 0
        self.full = False

    def add(self, line: str) -> bool:
        if self.full:
            return False
        if self.used + len(line) > MAX_TEXT_CHARS:
            self.lines.append(f"[truncated: text exceeded {MAX_TEXT_CHARS} characters]")
            self.full = True
            return False
        self.lines.append(line)
        self.used += len(line) + 1
        return True

    def text(self) -> str:
        return "\n".join(self.lines)


def _tsv(rows, budget: _TextBudget, max_rows: int | None = None) -> None:
    for index, row in enumerate(rows):
        if max_rows is not None and index >= max_rows:  # counts empty rows too (sparse-sheet bombs)
            budget.add(f"[truncated: sheet exceeded {max_rows} rows]")
            return
        cells = [_cell(v) for v in row]
        while cells and not cells[-1]:
            cells.pop()
        if any(cells) and not budget.add("\t".join(cells)):
            return


def _check_zip(data: bytes, what: str) -> None:
    """Reject OOXML zips that would inflate absurdly (zip bombs) before any parser touches them."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        raise ExtractionError(f"Could not open {what}: not a valid zip container ({exc})") from exc
    total = sum(i.file_size for i in infos)
    if total > MAX_ZIP_UNCOMPRESSED:
        raise ExtractionError(f"{what} is too large when uncompressed ({total // (1024 * 1024)} MB)")
    for i in infos:
        if i.file_size > _ZIP_RATIO_MIN_SIZE and i.file_size / max(i.compress_size, 1) > MAX_ZIP_RATIO:
            raise ExtractionError(f"{what} rejected: entry {i.filename!r} is suspiciously highly compressed")


def to_png(img, max_edge: int = MAX_IMAGE_EDGE) -> bytes:
    """PIL image -> PNG bytes, RGB/L, long edge downscaled to ``max_edge``."""
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    if max(img.size) > max_edge:
        img = img.copy()
        img.thumbnail((max_edge, max_edge), Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# per-kind extractors
# ---------------------------------------------------------------------------

def _extract_text(data: bytes, filename: str) -> ExtractedContent:
    text = data.decode("utf-8-sig", errors="replace")
    return ExtractedContent(text=text, pages=1)


def _extract_docx(data: bytes, filename: str) -> ExtractedContent:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _check_zip(data, "Word document")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - python-docx raises assorted errors on corrupt zips
        raise ExtractionError(f"Could not open Word document: {exc}") from exc
    budget = _TextBudget()
    for child in document.element.body.iterchildren():
        if budget.full:
            break
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = Paragraph(child, document).text.strip()
            if t:
                budget.add(t)
        elif tag == "tbl":
            table = Table(child, document)
            rows = []
            for row in table.rows:
                cells, prev = [], None
                for c in row.cells:  # merged cells repeat the same object; keep one
                    if c._tc is prev:
                        continue
                    prev = c._tc
                    cells.append(c.text)
                rows.append(cells)
            _tsv(rows, budget)
    return ExtractedContent(text=budget.text(), pages=1)


def _extract_doc(data: bytes, filename: str) -> ExtractedContent:
    """Legacy .doc: best-effort scrape of printable runs (UTF-16LE text pieces and 8-bit runs)."""
    runs: list[str] = []
    for m in re.finditer(rb"(?:[\x20-\x7e]\x00){4,}", data):
        runs.append(m.group().decode("utf-16-le", errors="ignore"))
    if not runs:
        for m in re.finditer(rb"[\x20-\x7e\r\n\t]{6,}", data):
            runs.append(m.group().decode("latin-1"))
    text = "\n".join(r.strip() for r in runs if r.strip())[:MAX_TEXT_CHARS]
    return ExtractedContent(text=text, pages=1,
                            meta={"note": "Legacy .doc: text recovered by best-effort scrape; layout lost"})


def _extract_xlsx(data: bytes, filename: str) -> ExtractedContent:
    import openpyxl

    _check_zip(data, "Excel workbook")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Could not open Excel workbook: {exc}") from exc
    budget = _TextBudget()
    names = []
    try:
        for ws in wb.worksheets:
            names.append(ws.title)
            if budget.add(f"[Sheet: {ws.title}]"):
                _tsv(ws.iter_rows(values_only=True), budget, max_rows=MAX_SHEET_ROWS)
    finally:
        wb.close()
    return ExtractedContent(text=budget.text(), pages=len(names), meta={"sheets": names})


def _extract_xls(data: bytes, filename: str) -> ExtractedContent:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Could not open Excel 97 workbook: {exc}") from exc
    budget = _TextBudget()
    names = []
    for sheet in book.sheets():
        names.append(sheet.name)
        if not budget.add(f"[Sheet: {sheet.name}]"):
            continue
        for r in range(sheet.nrows):
            if budget.full:
                break
            if r >= MAX_SHEET_ROWS:
                budget.add(f"[truncated: sheet exceeded {MAX_SHEET_ROWS} rows]")
                break
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                        continue
                    except Exception:  # noqa: BLE001
                        pass
                row.append(cell.value)
            _tsv([row], budget)
    return ExtractedContent(text=budget.text(), pages=len(names), meta={"sheets": names})


def _render_pdf_pages(data: bytes, page_numbers: list[int]) -> list[tuple[bytes, int]]:
    """Render pages to PNG. PDFium is not thread-safe, so the whole open/render/close sequence runs under
    ``_PDFIUM_LOCK`` (ingestion uses a thread pool)."""
    import pypdfium2 as pdfium

    images: list[tuple[bytes, int]] = []
    with _PDFIUM_LOCK:
        pdf = pdfium.PdfDocument(data)
        try:
            for n in page_numbers:
                page = pdf[n - 1]
                try:
                    w, h = page.get_size()  # points (1/72 in)
                    scale = min(200 / 72, MAX_IMAGE_EDGE / max(w, h, 1))  # ~200 dpi, capped to max edge
                    bitmap = page.render(scale=scale)
                    images.append((to_png(bitmap.to_pil()), n))
                finally:
                    page.close()
        finally:
            pdf.close()
    return images


def _extract_pdf(data: bytes, filename: str) -> ExtractedContent:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            reader.decrypt("")
        pages = reader.pages
        page_count = len(pages)
    except Exception as exc:  # noqa: BLE001 - pypdf raises a zoo of errors on malformed files
        raise ExtractionError(f"Could not read PDF: {exc}") from exc
    if page_count == 0:
        raise ExtractionError("PDF has no pages")
    parts: list[str] = []
    sparse: list[int] = []
    for i, page in enumerate(pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:  # noqa: BLE001 - one bad page shouldn't sink the document
            text = ""
        parts.append(f"[Page {i}]\n{text}" if text else f"[Page {i}]")
        if len(text) < MIN_PAGE_TEXT:
            sparse.append(i)
    to_render = sparse[:MAX_RENDERED_PAGES]
    images: list[tuple[bytes, int]] = []
    meta: dict = {"pdf_bytes": data, "page_count": page_count, "rendered_pages": []}
    if to_render:
        try:
            images = _render_pdf_pages(data, to_render)
            meta["rendered_pages"] = [n for _, n in images]
        except Exception as exc:  # noqa: BLE001
            log.warning("PDF render failed for %s: %s", filename, exc)
            meta["render_error"] = str(exc)
            if not any(len(p) > len(f"[Page {i}]") for i, p in enumerate(parts, 1)):
                raise ExtractionError(f"PDF has no text layer and could not be rendered: {exc}") from exc
        if len(sparse) > MAX_RENDERED_PAGES:
            meta["render_truncated"] = True
    return ExtractedContent(text="\n\n".join(parts), images=images, pages=page_count, meta=meta)


def _extract_image(data: bytes, filename: str) -> ExtractedContent:
    from PIL import ImageOps, UnidentifiedImageError

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            img = Image.open(io.BytesIO(data), formats=IMAGE_FORMATS)
            if img.format == "JPEG":
                img.draft("RGB", (MAX_IMAGE_EDGE * 2, MAX_IMAGE_EDGE * 2))  # decode at reduced scale
            img.load()
            img = ImageOps.exif_transpose(img)
    except (Image.DecompressionBombError, Image.DecompressionBombWarning) as exc:
        raise ExtractionError(f"Image is too large to process safely: {exc}") from exc
    except (UnidentifiedImageError, OSError, ValueError, SyntaxError) as exc:
        raise ExtractionError(f"Could not read image: {exc}") from exc
    original = img.size
    return ExtractedContent(images=[(to_png(img), 1)], pages=1, meta={"original_size": list(original)})


def _probe_duration(path: Path) -> float | None:
    ffprobe = shutil.which("ffprobe")
    if ffprobe:
        try:
            out = subprocess.run([ffprobe, "-v", "error", *_FFMPEG_SAFE, "-show_entries", "format=duration",
                                  "-of", "json", str(path)], capture_output=True, timeout=30, check=False)
            return float(json.loads(out.stdout or b"{}").get("format", {}).get("duration"))
        except (ValueError, TypeError, subprocess.SubprocessError, OSError):
            pass
    ffmpeg = shutil.which("ffmpeg")
    try:
        out = subprocess.run([ffmpeg, "-hide_banner", *_FFMPEG_SAFE, "-i", str(path)], capture_output=True,
                             timeout=30, check=False)
        m = re.search(rb"Duration:\s*(\d+):(\d+):(\d+(?:\.\d+)?)", out.stderr)
        if m:
            h, mi, s = m.groups()
            return int(h) * 3600 + int(mi) * 60 + float(s)
    except (subprocess.SubprocessError, OSError):
        pass
    return None


def _extract_video(data: bytes, filename: str) -> ExtractedContent:
    ffmpeg = shutil.which("ffmpeg")
    if ffmpeg is None:
        raise ExtractionError("Video files need ffmpeg, which is not installed on this server")

    suffix = Path(filename).suffix or ".mp4"
    with tempfile.TemporaryDirectory(prefix="idhv-video-") as tmp:
        src = Path(tmp) / f"input{suffix}"
        src.write_bytes(data)
        duration = _probe_duration(src)
        if not duration or duration <= 0:
            raise ExtractionError("Could not read video (no duration; file may be corrupt)")
        fps = MAX_VIDEO_FRAMES / duration
        cmd = [ffmpeg, "-hide_banner", "-loglevel", "error", "-y", *_FFMPEG_SAFE, "-i", str(src),
               "-vf", f"fps={fps:.6f},scale='min({MAX_IMAGE_EDGE},iw)':-2",
               "-frames:v", str(MAX_VIDEO_FRAMES), str(Path(tmp) / "frame_%02d.png")]
        try:
            proc = subprocess.run(cmd, capture_output=True, timeout=FFMPEG_TIMEOUT, check=False)
        except subprocess.TimeoutExpired as exc:
            raise ExtractionError("ffmpeg timed out extracting frames") from exc
        except OSError as exc:
            raise ExtractionError(f"ffmpeg could not be run: {exc}") from exc
        frames = sorted(Path(tmp).glob("frame_*.png"))
        if not frames:
            err = proc.stderr.decode(errors="replace").strip()[-300:]
            raise ExtractionError(f"ffmpeg could not extract frames: {err or 'no frames produced'}")
        images = []
        for i, f in enumerate(frames[:MAX_VIDEO_FRAMES], start=1):
            with Image.open(f, formats=["PNG"]) as im:
                images.append((to_png(im), i))
    times = [round((i - 0.5) / fps, 2) for i in range(1, len(images) + 1)]
    return ExtractedContent(images=images, pages=len(images),
                            meta={"duration_s": round(duration, 2), "frame_times_s": times})


_EXTRACTORS = {
    "text": _extract_text,
    "docx": _extract_docx,
    "doc": _extract_doc,
    "xlsx": _extract_xlsx,
    "xls": _extract_xls,
    "pdf": _extract_pdf,
    "image": _extract_image,
    "video": _extract_video,
}


def extract(kind: str, data: bytes, filename: str) -> ExtractedContent:
    """Extract text/images from ``data`` of ``kind``. Raises ``ExtractionError`` on unreadable input."""
    fn = _EXTRACTORS.get(kind)
    if fn is None:
        raise ExtractionError(f"Unsupported file type: {kind}")
    if not data:
        raise ExtractionError("File is empty")
    content = fn(data, filename)
    content.meta["kind"] = kind
    return content
