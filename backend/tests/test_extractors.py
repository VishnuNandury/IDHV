"""Task 2: kind detection + multi-format extraction (fixtures generated in tmp)."""
from __future__ import annotations

import io
import shutil
import subprocess
from pathlib import Path

import pytest
from PIL import Image

from app.ingestion.detect import ALLOWED_EXTENSIONS, detect_kind
from app.ingestion.extractors import ExtractionError, extract

SAMPLES = Path(__file__).resolve().parents[2] / "samples"


# ---- fixture builders --------------------------------------------------------

def _docx_bytes() -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("Discharge summary for synthetic patient")
    t = d.add_table(rows=2, cols=3)
    for i, v in enumerate(("Test", "Result", "Unit")):
        t.cell(0, i).text = v
    for i, v in enumerate(("HbA1c", "7.4", "%")):
        t.cell(1, i).text = v
    d.add_paragraph("Follow up after 2 weeks")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HbA1c"
    ws.append(["Date", "Value"])
    ws.append(["2024-01-18", 7.9])
    ws2 = wb.create_sheet("Lipids")
    ws2.append(["Date", "LDL"])
    ws2.append(["2024-02-10", 142])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _text_pdf_bytes() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.drawString(72, 750, "Sunrise Diagnostics lab report page one with enough text to count")
    c.drawString(72, 730, "HbA1c 7.0 % 4.0 - 5.6")
    c.showPage()
    c.drawString(72, 750, "Second page: Serum Creatinine 0.9 mg/dL 0.5 - 1.1 reference range")
    c.showPage()
    c.save()
    return buf.getvalue()


def _png_bytes(size=(400, 300), color=(200, 30, 30)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="PNG")
    return buf.getvalue()


def _image_only_pdf_bytes() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.drawImage(ImageReader(io.BytesIO(_png_bytes((600, 800)))), 50, 50, width=400, height=600)
    c.showPage()
    c.save()
    return buf.getvalue()


def _jpg_bytes(size=(3000, 2000)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, (240, 240, 230)).save(buf, format="JPEG")
    return buf.getvalue()


# ---- detect ------------------------------------------------------------------

def test_detect_kind_by_magic_and_extension():
    assert detect_kind("x.pdf", b"%PDF-1.7\n") == "pdf"
    assert detect_kind("x.docx", b"PK\x03\x04rest") == "docx"
    assert detect_kind("x.xlsx", b"PK\x03\x04rest") == "xlsx"
    assert detect_kind("x.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1") == "xls"
    assert detect_kind("x.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1") == "doc"
    assert detect_kind("photo.jpeg", b"\xff\xd8\xff\xe0") == "image"
    assert detect_kind("scan.png", b"\x89PNG\r\n\x1a\n") == "image"
    assert detect_kind("clip.mp4", b"\x00\x00\x00\x18ftypmp42") == "video"
    assert detect_kind("clip.mkv", b"\x1a\x45\xdf\xa3") == "video"
    assert detect_kind("notes.txt", b"hello") == "text"
    assert detect_kind("data.csv", b"a,b\n1,2") == "text"
    # magic bytes win over a misleading extension
    assert detect_kind("report.txt", b"%PDF-1.4") == "pdf"
    assert detect_kind("virus.exe", b"MZ\x90\x00") == "other"
    assert "exe" not in ALLOWED_EXTENSIONS and "mkv" in ALLOWED_EXTENSIONS


# ---- extract -----------------------------------------------------------------

def test_docx_paragraphs_and_tables():
    c = extract("docx", _docx_bytes(), "d.docx")
    assert "Discharge summary for synthetic patient" in c.text
    assert "HbA1c\t7.4\t%" in c.text  # table rows as TSV
    assert c.text.index("HbA1c") < c.text.index("Follow up")  # document order preserved
    assert c.meta["kind"] == "docx"


def test_xlsx_two_sheets():
    c = extract("xlsx", _xlsx_bytes(), "h.xlsx")
    for needle in ("HbA1c", "Lipids", "7.9", "142", "2024-02-10"):
        assert needle in c.text
    assert c.pages == 2
    assert c.meta["kind"] == "xlsx"


def test_text_pdf_page_markers():
    data = _text_pdf_bytes()
    c = extract("pdf", data, "lab.pdf")
    assert "[Page 1]" in c.text and "[Page 2]" in c.text
    assert "HbA1c 7.0" in c.text
    assert c.text.index("[Page 2]") < c.text.index("Creatinine")
    assert c.images == []  # pages have text, nothing rendered
    assert c.meta["pdf_bytes"] == data
    assert c.meta["page_count"] == 2 and c.pages == 2
    assert c.meta["kind"] == "pdf"


def test_image_only_pdf_renders_pages():
    c = extract("pdf", _image_only_pdf_bytes(), "scan.pdf")
    assert len(c.images) == 1
    png, page = c.images[0]
    assert page == 1 and png.startswith(b"\x89PNG")
    assert max(Image.open(io.BytesIO(png)).size) <= 1568
    assert c.meta["rendered_pages"] == [1]


def test_jpg_downscaled_to_png():
    c = extract("image", _jpg_bytes(), "rx.jpg")
    assert len(c.images) == 1
    png, page = c.images[0]
    assert png.startswith(b"\x89PNG")
    im = Image.open(io.BytesIO(png))
    assert max(im.size) == 1568
    assert c.meta["kind"] == "image"


def test_text_decode_with_replacement():
    c = extract("text", "BP 120/80\n".encode() + b"\xff", "n.txt")
    assert "BP 120/80" in c.text


def test_legacy_doc_best_effort_scrape():
    blob = b"\xd0\xcf\x11\xe0" + b"\x00" * 20 + "Tab Metformin 500 mg BD".encode("utf-16-le") + b"\x00\x01" * 5
    c = extract("doc", blob, "old.doc")
    assert "Tab Metformin 500 mg BD" in c.text
    assert "note" in c.meta


def test_corrupt_files_raise_extraction_error():
    with pytest.raises(ExtractionError):
        extract("pdf", b"%PDF-garbage", "bad.pdf")
    with pytest.raises(ExtractionError):
        extract("image", b"not an image", "bad.jpg")
    with pytest.raises(ExtractionError):
        extract("other", b"MZ", "x.exe")


@pytest.mark.skipif(shutil.which("ffmpeg") is None, reason="ffmpeg not installed")
def test_video_frames(tmp_path):
    out = tmp_path / "v.mp4"
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=duration=4:size=320x240:rate=10",
                    "-pix_fmt", "yuv420p", str(out)], check=True)
    c = extract("video", out.read_bytes(), "v.mp4")
    assert 1 <= len(c.images) <= 6
    assert all(p.startswith(b"\x89PNG") for p, _ in c.images)
    assert c.meta["kind"] == "video"


def test_video_without_ffmpeg_fails_clearly(monkeypatch):
    import app.ingestion.extractors as ex

    monkeypatch.setattr(ex.shutil, "which", lambda name: None)
    with pytest.raises(ExtractionError, match="ffmpeg"):
        extract("video", b"\x00\x00\x00\x18ftypmp42", "v.mp4")


@pytest.mark.parametrize("path", sorted(p for p in SAMPLES.glob("*") if p.suffix.lower() not in (".md", ".py")),
                         ids=lambda p: p.name)
def test_samples_smoke(path):
    data = path.read_bytes()
    kind = detect_kind(path.name, data[:16])
    if kind == "video" and shutil.which("ffmpeg") is None:
        pytest.skip("ffmpeg not installed")
    c = extract(kind, data, path.name)
    assert c.text.strip() or c.images, f"{path.name} produced nothing"


# ---- fix round 1: hostile inputs ----------------------------------------------

def test_zip_bomb_docx_rejected():
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", b"\x00" * (5 * 1024 * 1024))  # ~1000:1
    with pytest.raises(ExtractionError, match="compress"):
        extract("docx", buf.getvalue(), "bomb.docx")
    with pytest.raises(ExtractionError, match="compress"):
        extract("xlsx", buf.getvalue(), "bomb.xlsx")


def test_zip_total_size_limit(monkeypatch):
    import app.ingestion.extractors as ex

    monkeypatch.setattr(ex, "MAX_ZIP_UNCOMPRESSED", 1000)
    with pytest.raises(ExtractionError, match="too large"):
        extract("docx", _docx_bytes(), "d.docx")


def test_tsv_output_is_capped(monkeypatch):
    import app.ingestion.extractors as ex

    monkeypatch.setattr(ex, "MAX_TEXT_CHARS", 20)
    c = extract("xlsx", _xlsx_bytes(), "h.xlsx")
    assert "[truncated" in c.text and len(c.text) < 120
    c = extract("docx", _docx_bytes(), "d.docx")
    assert "[truncated" in c.text


def test_decompression_bomb_image_rejected(monkeypatch):
    monkeypatch.setattr(Image, "MAX_IMAGE_PIXELS", 1000)
    with pytest.raises(ExtractionError, match="too large"):
        extract("image", _png_bytes((100, 100)), "bomb.png")


def test_image_pixel_limit_default():
    import app.ingestion.extractors  # noqa: F401

    assert Image.MAX_IMAGE_PIXELS == 100_000_000  # 48-50 MP phone photos must decode


def test_image_formats_restricted():
    buf = io.BytesIO()
    Image.new("RGB", (10, 10)).save(buf, format="GIF")
    with pytest.raises(ExtractionError):
        extract("image", buf.getvalue(), "x.png")


def test_extension_magic_mismatch_for_media():
    assert detect_kind("photo.jpg", b"hello this is text") == "other"
    assert detect_kind("clip.mp4", b"#EXTM3U\n#EXTINF:-1,\nhttp://evil/x") == "other"
    assert detect_kind("photo.jpg", b"%PDF-1.4") == "other"
    assert detect_kind("photo.png", b"\xff\xd8\xff\xe0") == "image"  # still an image, just misnamed


def test_sniff_mime():
    from app.ingestion.detect import sniff_mime

    assert sniff_mime(b"\x89PNG\r\n\x1a\n") == "image/png"
    assert sniff_mime(b"\xff\xd8\xff\xe0") == "image/jpeg"
    assert sniff_mime(b"%PDF-1.4") == "application/pdf"
    assert sniff_mime(b"\x00\x00\x00\x14ftypqt  ") == "video/quicktime"
    assert sniff_mime(b"hello") is None


def test_ffmpeg_protocol_whitelist(monkeypatch):
    import app.ingestion.extractors as ex

    calls = []

    def fake_run(cmd, **kw):
        calls.append(cmd)
        raise OSError("stop")

    monkeypatch.setattr(ex.shutil, "which", lambda n: "/usr/bin/" + n)
    monkeypatch.setattr(ex.subprocess, "run", fake_run)
    with pytest.raises(ExtractionError):
        extract("video", b"\x00\x00\x00\x18ftypmp42", "v.mp4")
    assert calls and all("-protocol_whitelist" in c and c[c.index("-protocol_whitelist") + 1] == "file"
                         for c in calls)


# ---- fix round 2 ---------------------------------------------------------------

def test_sheet_row_cap_counts_empty_rows(monkeypatch):
    import openpyxl

    import app.ingestion.extractors as ex

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Big"
    ws.append(["header"])
    ws.append([None])  # empty row still counts
    ws.append(["row three"])
    buf = io.BytesIO()
    wb.save(buf)
    monkeypatch.setattr(ex, "MAX_SHEET_ROWS", 2)
    c = extract("xlsx", buf.getvalue(), "big.xlsx")
    assert "header" in c.text and "row three" not in c.text
    assert "[truncated" in c.text and "Big" in c.text
